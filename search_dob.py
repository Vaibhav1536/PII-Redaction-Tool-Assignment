import docx
import re

doc_path = r"c:\Users\hp\OneDrive\Desktop\ksh_project\Red Herring Prospectus.docx"
doc = docx.Document(doc_path)

keywords = ["birth", "born", "age", "date of"]
matches = []

for idx, p in enumerate(doc.paragraphs):
    text = p.text.lower()
    for kw in keywords:
        if kw in text:
            matches.append((f"p[{idx}]", p.text))
            break

for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            text = cell.text.lower()
            for kw in keywords:
                if kw in text:
                    matches.append((f"t[{t_idx}]r[{r_idx}]c[{c_idx}]", cell.text))
                    break

print(f"Total keyword matches: {len(matches)}")
for ref, text in matches[:20]:
    print(f"  {ref}: {text[:100]}...")
