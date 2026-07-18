import docx
import json
import re

doc_path = r"c:\Users\hp\OneDrive\Desktop\ksh_project\Red Herring Prospectus.docx"
doc = docx.Document(doc_path)

# Ground Truth Lists (Complete & Verified)
true_emails = [
    "cs.connect@kshinternational.com",
    "Sarthak.malvadkar@kshinterantional.com",
    "ksh.ipo@nuvama.com",
    "customerservice.mb@nuvama.com",
    "ksh@icicisecurities.com",
    "customercare@icicisecurities.com",
    "prakash.boricha@nuvama.com",
    "sheetal.parab@nuvama.com",
    "ipo@trilegal.com",
    "kshinternational.ipo@in.mpms.mufg.com",
    "siddharth.jadhav@hdfcbank.com",
    "sachin.gawade@hdfcbank.com",
    "eric.bacha@hdfcbank.com",
    "tushar.gavankar@hdfcbank.com",
    "pravin.teli2@hdfcbank.com",
    "Ipocmg@icicibank.com",
    # Newly discovered real emails
    "parag.pansare@kirtanepandit.com",
    "pro@eximbankindia.in",
    "hitesh.ramani@citi.com",
    "anand.soni@bajajfinserv.in",
    "ashishmp@federalbank.co.in",
    "rm6.ifbpune@sbi.co.in",
    "sharmila.joshi@indusind.com",
    "cherag.gyara@icicibank.com",
    "manisha.shukla@hdfcbank.com",
    "hingnetare@gmail.com"
]

true_phones = [
    "+ 91 20 4505 3237",
    "+ 91 20 45053237",
    "+91 22 40094400",
    "+91 22 6807 7100",
    "+ 91 22 4009 4400",
    "+91 22 4079 1000",
    "+91 81081 14949",
    "+91 22 30752929",
    "+91 22 30752928",
    "+91 22 30752914",
    "+91 20 6606 4494",
    "+91 20 2640 3100",
    "+ 91 8879770456",
    # Spacing variations
    "+91 20 4505 3237",
    "+91 20 45053237",
    "+ 91 22 40094400",
    "+ 91 22 6807 7100",
    "+91 22 4009 4400",
    "+ 91 22 4079 1000",
    "+ 91 81081 14949",
    "+ 91 22 30752929",
    "+ 91 22 30752928",
    "+ 91 22 30752914",
    "+ 91 20 6606 4494",
    "+ 91 20 2640 3100",
    "+91 8879770456"
]

true_names = [
    "Sarthak Malvadkar",
    "Kushal Subbayya Hegde",
    "Kushal Hegde",
    "Rajesh Kushal Hegde",
    "Rajesh Hegde",
    "Rohit Kushal Hegde",
    "Rohit Hegde",
    "Pushpa Kushal Hegde",
    "Pushpa Hegde",
    "Rakhi Girija Shetty",
    "Sandesh Bhagwat",
    "Amod Joshi",
    "Ganesh Prasad",
    "Dinesh Hirachand Munot",
    "Ajay Shriram Patil",
    "Kishan Rastogi",
    "Abhijit Diwan",
    "Shanti Gopalkrishnan",
    "Ajay Menon",
    "Karunakar Hegde",
    "Vijay Hegde",
    "Karunakar N. Bhandary",
    "Narayna B. Shetty",
    "Jayaram N. Shetty",
    "Sangeeta Ramprasad Rai",
    "Sharmila Joshi",
    "Cherag Gyara",
    "Rupal K. Sancheti",
    "Salil Ajay Bhargava",
    "Indu Jacob",
    "Kumar Tiwari",
    "Siddharth Jadhav",
    "Sachin Gawade",
    "Eric Bacha",
    "Tushar Gavankar",
    "Pravin Teli",
    "Prakash Boricha",
    "Sheetal Parab",
    # Newly discovered names
    "Parag Pansare",
    "Hitesh Ramani",
    "Anand Soni",
    "Manisha Shukla",
    "Ashish MP"
]

all_name_variations = set()
for name in true_names:
    all_name_variations.add(name)
    all_name_variations.add(name.upper())
    all_name_variations.add(name.lower())
    parts = name.split()
    if len(parts) == 3:
        short_name = f"{parts[0]} {parts[2]}"
        all_name_variations.add(short_name)
        all_name_variations.add(short_name.upper())
        all_name_variations.add(short_name.lower())

true_companies = [
    "KSH INTERNATIONAL LIMITED",
    "KSH International Limited",
    "Bhandary Metal Extrusion Private Limited",
    "KSH International Private Limited",
    "KSH International",
    "Nuvama Wealth Management Limited",
    "Nuvama",
    "ICICI Securities Limited",
    "ICICI Securities",
    "Link Intime India Private Limited",
    "Link Intime",
    "CareEdge Research",
    "CARE Analytics and Advisory Private Limited",
    "CARE Analytics",
    "Kanj and Co LLP",
    "Kanj & Co",
    "Trilegal",
    "HDFC Bank",
    "ICICI Bank",
    "Waterloo Industrial Park VI Private Limited",
    "Waterloo Industrial Park III Private Limited",
    "Waterloo Industrial Park IV Private Limited",
    "Waterloo Industrial Park V Private Limited",
    "Waterloo Industrial Park II Private Limited",
    "Waterloo Industrial Park VIII Private Limited",
    "Waterloo Industrial Park IX Private Limited",
    "Waterloo Industrial Park IX B Private Limited",
    "Waterloo Industrial Park",
    "Waterloo Industrial",
    "Bijlee Limited",
    "Nidec Industrial Automation India Private Limited",
    "Kushal Electricals",
    # Newly discovered companies
    "Kirtane & Pandit LLP",
    "Export-Import Bank of India",
    "Citi",
    "Citibank",
    "Bajaj Finserv",
    "Federal Bank",
    "State Bank of India",
    "SBI",
    "IndusInd Bank"
]

true_addresses = [
    "11/3, 11/4 and 11/5, Village Birdewadi, Chakan Taluka - Khed, Pune – 410 501, Maharashtra, India",
    "201, Tower 2, Montreal Business Centre, Off Pallod Farms, Baner, Pune – 411 045, Maharashtra, India",
    "Unit 3 located at Chakan, Pune, Maharashtra",
    "Plot No. F-223, Supa Parner Industrial Park, Mauje Palve Khurd, Taluka Parner, Dist – Ahmednagar, Maharashtra – 414 301",
    "Taloja (Raigad), Maharashtra",
    "Chakan (Pune), Maharashtra",
    "Supa, Ahilyanagar (formerly Ahmednagar) in Maharashtra",
    "Village Khalumbre, Chakan Taluka-Khed, Pune",
    "Plot No. F-223, Supa Parner Industrial Park, Mauje Palve Khurd, Taluka Parner, Dist – Ahmednagar, Maharashtra – 414 301"
]

# Compile Combined Regexes
sorted_names = sorted(list(all_name_variations), key=len, reverse=True)
sorted_companies = sorted(true_companies, key=len, reverse=True)
sorted_addresses = sorted(true_addresses, key=len, reverse=True)

email_regex = re.compile(r'\b(' + '|'.join(map(re.escape, sorted(true_emails, key=len, reverse=True))) + r')\b', re.I)
phone_regex = re.compile(r'(' + '|'.join(map(re.escape, sorted(true_phones, key=len, reverse=True))) + r')')
name_regex = re.compile(r'\b(' + '|'.join(map(re.escape, sorted_names)) + r')\b')
company_regex = re.compile(r'\b(' + '|'.join(map(re.escape, sorted_companies)) + r')\b')
address_prefix_regex = re.compile(r'(' + '|'.join(map(re.escape, [ad[:30] for ad in sorted_addresses])) + r')', re.I)

email_counts = 0
phone_counts = 0
name_counts = 0
company_counts = 0
address_counts = 0

def count_in_text(text):
    global email_counts, phone_counts, name_counts, company_counts, address_counts
    email_counts += len(email_regex.findall(text))
    phone_counts += len(phone_regex.findall(text))
    name_counts += len(name_regex.findall(text))
    company_counts += len(company_regex.findall(text))
    address_counts += len(address_prefix_regex.findall(text))

# Process doc paragraphs (tracking processed elements to avoid duplicates from merged cells)
processed_paragraphs = set()

for p in doc.paragraphs:
    p_id = id(p._element)
    if p_id in processed_paragraphs:
        continue
    processed_paragraphs.add(p_id)
    if p.text.strip():
        count_in_text(p.text)

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p_id = id(p._element)
                if p_id in processed_paragraphs:
                    continue
                processed_paragraphs.add(p_id)
                if p.text.strip():
                    count_in_text(p.text)

print("Counts compiled:")
print(f"  Emails: {email_counts}")
print(f"  Phones: {phone_counts}")
print(f"  Names: {name_counts}")
print(f"  Companies: {company_counts}")
print(f"  Addresses: {address_counts}")

gt_data = {
    "counts": {
        "email": email_counts,
        "phone": phone_counts,
        "name": name_counts,
        "company": company_counts,
        "address": address_counts,
        "ssn": 0,
        "credit_card": 0,
        "dob": 0,
        "ip": 0
    },
    "entities": {
        "emails": true_emails,
        "phones": true_phones,
        "names": true_names,
        "companies": true_companies,
        "addresses": true_addresses
    }
}

gt_path = r"c:\Users\hp\OneDrive\Desktop\ksh_project\ground_truth.json"
with open(gt_path, "w", encoding="utf-8") as f:
    json.dump(gt_data, f, indent=4)

print("ground_truth.json saved successfully!")
