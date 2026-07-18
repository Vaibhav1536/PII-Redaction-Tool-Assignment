import docx
import json
import re
import spacy
import os
import random

# Seed for reproducibility of random elements
random.seed(42)

# Load spaCy
nlp = spacy.load("en_core_web_sm")

# Define paths
doc_path = r"c:\Users\hp\OneDrive\Desktop\ksh_project\Red Herring Prospectus.docx"
output_path = r"c:\Users\hp\OneDrive\Desktop\ksh_project\Red Herring Prospectus_redacted.docx"
detections_path = r"c:\Users\hp\OneDrive\Desktop\ksh_project\detections.json"

# List of known names (Gazetteer)
known_names = [
    "Sarthak Malvadkar", "Kushal Subbayya Hegde", "Kushal Hegde", 
    "Rajesh Kushal Hegde", "Rajesh Hegde", "Rohit Kushal Hegde", "Rohit Hegde", 
    "Pushpa Kushal Hegde", "Pushpa Hegde", "Rakhi Girija Shetty", 
    "Sandesh Bhagwat", "Amod Joshi", "Ganesh Prasad", 
    "Dinesh Hirachand Munot", "Ajay Shriram Patil", "Kishan Rastogi", 
    "Abhijit Diwan", "Shanti Gopalkrishnan", "Ajay Menon", "Karunakar Hegde", 
    "Vijay Hegde", "Karunakar N. Bhandary", "Narayna B. Shetty", 
    "Jayaram N. Shetty", "Sangeeta Ramprasad Rai", "Sharmila Joshi", 
    "Cherag Gyara", "Rupal K. Sancheti", "Salil Ajay Bhargava", 
    "Indu Jacob", "Kumar Tiwari", "Siddharth Jadhav", "Sachin Gawade", 
    "Eric Bacha", "Tushar Gavankar", "Pravin Teli", "Prakash Boricha", 
    "Sheetal Parab", "Parag Pansare", "Hitesh Ramani", "Anand Soni", 
    "Manisha Shukla", "Ashish MP"
]

# Name blacklist for spaCy PERSON false positives
blacklist_names = {
    "offer", "directors", "promoters", "company", "app", "sebi", "brlms", 
    "stock", "exchange", "equity", "shares", "bidders", "bidder", "bid", 
    "bids", "members", "member", "auditors", "auditor", "legal", "advisors", 
    "advisor", "managers", "manager", "lead", "running", "book", "syndicate", 
    "registrar", "compliances", "compliance", "secretary", "officer", 
    "officers", "promoter", "director", "companies", "act", "rules", 
    "regulations", "prospectus", "issue", "allotment", "anchor", "investors", 
    "investor", "individuals", "individual", "non-institutional", 
    "institutional", "category", "net", "proceeds", "objects", 
    "manufacturing", "facilities", "facility", "unit", "units", "plant", 
    "plants", "land", "lease", "leases", "date", "dates", "year", "years", 
    "fiscals", "fiscal", "crore", "million", "lakh", "rupees", "rs.", "inr", 
    "percent", "percentage", "average", "total", "sum", "table", "page", 
    "pages", "schedule", "section", "annexure", "mumbai", "pune", 
    "maharashtra", "india", "chakan", "supa", "taloja", "raigad", "ahmednagar", 
    "parner", "care", "letter", "gaap", "cagr", "ebitda", "net proceeds", 
    "outstanding litigation", "history", "certain corporate", "summary", 
    "capital structure", "risk factors", "objects of the offer", "basis", 
    "offer price"
}

# Known Companies
known_companies = [
    "KSH INTERNATIONAL LIMITED", "KSH International Limited", 
    "Bhandary Metal Extrusion Private Limited", "KSH International Private Limited", 
    "KSH International", "Nuvama Wealth Management Limited", "Nuvama", 
    "ICICI Securities Limited", "ICICI Securities", 
    "Link Intime India Private Limited", "Link Intime", 
    "CareEdge Research", "CARE Analytics and Advisory Private Limited", 
    "CARE Analytics", "Kanj and Co LLP", "Kanj & Co", "Trilegal", 
    "HDFC Bank", "ICICI Bank", "Bijlee Limited", 
    "Nidec Industrial Automation India Private Limited", "Kushal Electricals",
    "Waterloo Industrial Park VI Private Limited", "Waterloo Industrial Park III Private Limited", 
    "Waterloo Industrial Park IV Private Limited", "Waterloo Industrial Park V Private Limited", 
    "Waterloo Industrial Park II Private Limited", "Waterloo Industrial Park VIII Private Limited", 
    "Waterloo Industrial Park IX Private Limited", "Waterloo Industrial Park IX B Private Limited", 
    "Waterloo Industrial Park", "Waterloo Industrial",
    "Kirtane & Pandit LLP", "Export-Import Bank of India", "Citi", "Citibank", 
    "Bajaj Finserv", "Federal Bank", "State Bank of India", "SBI", "IndusInd Bank"
]

# Known Addresses
known_addresses = [
    "11/3, 11/4 and 11/5, Village Birdewadi, Chakan Taluka - Khed, Pune – 410 501, Maharashtra, India",
    "201, Tower 2, Montreal Business Centre, Off Pallod Farms, Baner, Pune – 411 045, Maharashtra, India",
    "Unit 3 located at Chakan, Pune, Maharashtra",
    "Plot No. F-223, Supa Parner Industrial Park, Mauje Palve Khurd, Taluka Parner, Dist – Ahmednagar, Maharashtra – 414 301",
    "Taloja (Raigad), Maharashtra",
    "Chakan (Pune), Maharashtra",
    "Supa, Ahilyanagar (formerly Ahmednagar) in Maharashtra",
    "Village Khalumbre, Chakan Taluka-Khed, Pune",
    "Plot No. F-223, Supa Parner Industrial Park, Mauje Palve Khurd, Taluka Parner, Dist – Ahmednagar, Maharashtra – 414 301"
]

# --- FAKE DATA GENERATOR AND CONSISTENT MAPPING ---
first_names_male = ["Aarav", "Vihaan", "Aditya", "Rohan", "Vikram", "Rahul", "Amit", "Sanjay", "Sandeep", "Deepak", "Rajesh", "Karan", "Vijay", "Anil", "Sunil", "Manish", "Ravi", "Alok"]
first_names_female = ["Pooja", "Neha", "Priya", "Anjali", "Kiran", "Sunita", "Ritu", "Deepika", "Shalini", "Rashmi", "Aditi", "Meera", "Swati", "Preeti", "Aisha", "Divya"]
last_names = ["Sharma", "Verma", "Gupta", "Rao", "Shetty", "Nair", "Joshi", "Patel", "Singh", "Kumar", "Mehta", "Desai", "Kulkarni", "Jadhav", "Shinde", "Sen", "Roy", "Das"]

surname_mapping = {
    "hegde": "Rao", "patil": "Joshi", "shetty": "Singhal", "malvadkar": "Sharma", 
    "joshi": "Kulkarni", "bhagwat": "Deshmukh", "prasad": "Mehta", "munot": "Chawla", 
    "rastogi": "Saxena", "diwan": "Trivedi", "gopalkrishnan": "Iyer", "menon": "Nair", 
    "bhandary": "Kamath", "rai": "Sinha", "sancheti": "Ranka", "bhargava": "Agrawal", 
    "jacob": "Abraham", "tiwari": "Mishra", "jadhav": "Shinde", "gawade": "Kadam", 
    "bacha": "D'Souza", "gavankar": "Sawant", "teli": "Patel", "boricha": "Solanki", 
    "parab": "Sawant", "pansare": "Patel", "ramani": "Rawat", "soni": "Sharma", 
    "shukla": "Sen", "mp": "Singh"
}

domain_mapping = {
    "kshinternational.com": "vanguardind.com",
    "kshinterantional.com": "vanguardind.com",
    "nuvama.com": "horizonwealth.com",
    "icicisecurities.com": "apexcap.com",
    "trilegal.com": "apexlegal.com",
    "hdfcbank.com": "indusbank.com",
    "icicibank.com": "indusbank.com",
    "in.mpms.mufg.com": "in.mpms.fakeg.com",
    "kirtanepandit.com": "beaconassurance.com",
    "eximbankindia.in": "statefinbank.in",
    "citi.com": "fakebank.com",
    "bajajfinserv.in": "horizonfinance.in",
    "federalbank.co.in": "indusbank.co.in",
    "sbi.co.in": "fakebank.co.in",
    "indusind.com": "fakebank.com"
}

name_map = {}
email_map = {}
phone_map = {}
company_map = {}
address_map = {}

# Prepopulate known maps for high consistency
predefined_names = {
    "Sarthak Malvadkar": "Rohan Sharma",
    "Kushal Subbayya Hegde": "Vikram Aditya Rao",
    "Kushal Hegde": "Vikram Rao",
    "Rajesh Kushal Hegde": "Aarav Vikram Rao",
    "Rajesh Hegde": "Aarav Rao",
    "Rohit Kushal Hegde": "Kabir Vikram Rao",
    "Rohit Hegde": "Kabir Rao",
    "Pushpa Kushal Hegde": "Sunita Vikram Rao",
    "Pushpa Hegde": "Sunita Rao",
    "Rakhi Girija Shetty": "Pooja Singhal",
    "Sandesh Bhagwat": "Sanjay Deshmukh",
    "Amod Joshi": "Abhishek Kulkarni",
    "Ganesh Prasad": "Vijay Mehta",
    "Dinesh Hirachand Munot": "Devendra Chawla",
    "Ajay Shriram Patil": "Animesh Joshi",
    "Parag Pansare": "Prakash Patel",
    "Hitesh Ramani": "Harish Rawat",
    "Anand Soni": "Arjun Sharma",
    "Manisha Shukla": "Meena Sen",
    "Ashish MP": "Amit Singh"
}

for k, v in predefined_names.items():
    name_map[k] = v
    name_map[k.upper()] = v.upper()
    name_map[k.lower()] = v.lower()

predefined_companies = {
    "KSH INTERNATIONAL LIMITED": "VANGUARD INDUSTRIES LIMITED",
    "KSH International Limited": "Vanguard Industries Limited",
    "Bhandary Metal Extrusion Private Limited": "Premier Metal Extrusions Private Limited",
    "KSH International Private Limited": "Vanguard Industries Private Limited",
    "KSH International": "Vanguard Industries",
    "Nuvama Wealth Management Limited": "Horizon Financial Advisory Limited",
    "Nuvama": "Horizon",
    "ICICI Securities Limited": "Apex Capital Markets Limited",
    "ICICI Securities": "Apex Securities",
    "Link Intime India Private Limited": "Secure Registry India Private Limited",
    "Link Intime": "Secure Registry",
    "CareEdge Research": "Beacon Market Research",
    "CARE Analytics and Advisory Private Limited": "Beacon Analytics and Advisory Private Limited",
    "CARE Analytics": "Beacon Analytics",
    "Kanj and Co LLP": "Legacy Audit & Assurance LLP",
    "Kanj & Co": "Legacy Audit",
    "Trilegal": "Apex Partners Legal",
    "HDFC Bank": "Indus Bank",
    "ICICI Bank": "Indus Bank",
    "Waterloo Industrial Park": "Stellar Industrial Park",
    "Waterloo Industrial": "Stellar Industrial",
    "Bijlee Limited": "Urja Power Limited",
    "Nidec Industrial Automation India Private Limited": "Precision Automation India Private Limited",
    "Kushal Electricals": "Power Electricals",
    "Kirtane & Pandit LLP": "Beacon Assurance & Advisory LLP",
    "Export-Import Bank of India": "State Development Bank of India",
    "Citi": "Indus Bank",
    "Citibank": "Indus Bank",
    "Bajaj Finserv": "Horizon Finance",
    "Federal Bank": "Vanguard Trust Bank",
    "State Bank of India": "National Bank of India",
    "SBI": "NBI",
    "IndusInd Bank": "Zenith Bank"
}

for k, v in predefined_companies.items():
    company_map[k] = v
    company_map[k.upper()] = v.upper()
    company_map[k.lower()] = v.lower()

predefined_addresses = {
    "11/3, 11/4 and 11/5, Village Birdewadi, Chakan Taluka - Khed, Pune – 410 501, Maharashtra, India": 
        "45/A, 45/B and 45/C, GIDC Industrial Estate, Taluka Gandhinagar, Ahmedabad – 382 010, Gujarat, India",
    "201, Tower 2, Montreal Business Centre, Off Pallod Farms, Baner, Pune – 411 045, Maharashtra, India": 
        "701, Block C, Silver Business Hub, Off Ring Road, Salt Lake, Kolkata – 700 091, West Bengal, India",
    "Unit 3 located at Chakan, Pune, Maharashtra": 
        "Unit 5 located at GIDC, Vadodara, Gujarat",
    "Plot No. F-223, Supa Parner Industrial Park, Mauje Palve Khurd, Taluka Parner, Dist – Ahmednagar, Maharashtra – 414 301": 
        "Plot No. C-99, Industrial Hub, Village Green, Taluka Green, Dist - Surat, Gujarat - 395003",
    "Taloja (Raigad), Maharashtra": "Morbi (Rajkot), Gujarat",
    "Chakan (Pune), Maharashtra": "Sanand (Ahmedabad), Gujarat",
    "Supa, Ahilyanagar (formerly Ahmednagar) in Maharashtra": "Vapi, Valsad in Gujarat",
    "Village Khalumbre, Chakan Taluka-Khed, Pune": "Village GIDC, Taluka Sanand, Ahmedabad"
}

for k, v in predefined_addresses.items():
    address_map[k] = v
    address_map[k.upper()] = v.upper()
    address_map[k.lower()] = v.lower()

# --- HELPER GENERATORS ---

def get_fake_name(real_name):
    if real_name in name_map:
        return name_map[real_name]
    
    clean_n = re.sub(r'^(Mr\.|Ms\.|Mrs\.|Shri|Smt\.|Dr\.)\s+', '', real_name, flags=re.I).strip()
    parts = clean_n.split()
    
    fake_parts = []
    if len(parts) == 0:
        return "John Doe"
        
    for i, part in enumerate(parts):
        part_lower = part.lower()
        if i == len(parts) - 1: # Last name (Surname)
            if part_lower in surname_mapping:
                fake_parts.append(surname_mapping[part_lower])
            else:
                chosen_surname = random.choice(last_names)
                surname_mapping[part_lower] = chosen_surname
                fake_parts.append(chosen_surname)
        else: # First or middle name
            if random.random() > 0.3:
                fake_parts.append(random.choice(first_names_male))
            else:
                fake_parts.append(random.choice(first_names_female))
                
    fake_name = " ".join(fake_parts)
    if real_name.isupper():
        fake_name = fake_name.upper()
    elif real_name.islower():
        fake_name = fake_name.lower()
        
    name_map[real_name] = fake_name
    return fake_name

def get_fake_email(real_email):
    if real_email in email_map:
        return email_map[real_email]
        
    parts = real_email.split('@')
    if len(parts) != 2:
        return "redacted@example.com"
    local, domain = parts[0], parts[1]
    
    domain_lower = domain.lower()
    fake_domain = domain_mapping.get(domain_lower, "example.com")
    
    fake_local = local
    local_lower = local.lower()
    matched_name = False
    for real, fake in name_map.items():
        real_parts = real.lower().split()
        if len(real_parts) >= 2:
            if real_parts[0] in local_lower and real_parts[-1] in local_lower:
                fake_parts = fake.lower().split()
                fake_local = f"{fake_parts[0]}.{fake_parts[-1]}"
                matched_name = True
                break
                
    if not matched_name:
        fake_local = "".join(random.choice('abcdefghijklmnopqrstuvwxyz') for _ in range(8))
        
    fake_email = f"{fake_local}@{fake_domain}"
    if real_email.isupper():
        fake_email = fake_email.upper()
        
    email_map[real_email] = fake_email
    return fake_email

def get_fake_phone(real_phone):
    if real_phone in phone_map:
        return phone_map[real_phone]
        
    digits = re.sub(r'\D', '', real_phone)
    if digits.startswith('91') and len(digits) > 10:
        prefix = "+91 "
        if '20' in real_phone:
            prefix += "20 "
        elif '22' in real_phone:
            prefix += "22 "
        
        suffix = "".join(str(random.randint(0, 9)) for _ in range(8 if "20" in prefix or "22" in prefix else 10))
        fake_phone = prefix + suffix
    else:
        fake_phone = re.sub(r'\d', lambda x: str(random.randint(0, 9)), real_phone)
        
    phone_map[real_phone] = fake_phone
    return fake_phone

def get_fake_company(real_company):
    if real_company in company_map:
        return company_map[real_company]
        
    prefix = random.choice(["Apex", "Prime", "Global", "Summit", "Sterling", "Horizon", "Vanguard", "Delta", "Nexus"])
    sector = random.choice(["Solutions", "Industries", "Technologies", "Holdings", "Enterprises", "Capital", "Ventures"])
    suffix = ""
    if "Limited" in real_company or "LIMITED" in real_company:
        suffix = " Limited"
    elif "Pvt. Ltd." in real_company or "Private Limited" in real_company:
        suffix = " Private Limited"
    elif "LLP" in real_company:
        suffix = " LLP"
        
    fake_company = f"{prefix} {sector}{suffix}"
    if real_company.isupper():
        fake_company = fake_company.upper()
        
    company_map[real_company] = fake_company
    return fake_company

def get_fake_address(real_address):
    if real_address in address_map:
        return address_map[real_address]
        
    fake_address = f"{random.randint(10, 199)}/A, Industrial Sector {random.randint(1, 10)}, GIDC Estate, Sanand, Ahmedabad - 382 110, Gujarat, India"
    address_map[real_address] = fake_address
    return fake_address

# --- DETECTION ENGINE ---

email_pat = re.compile(r'\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}\b')
phone_pat = re.compile(r'\+?\s*\d{1,4}(?:\s*\d{2,5}){2,4}')
ssn_pat = re.compile(r'\b\d{3}-\d{2}-\d{4}\b')
cc_pat = re.compile(r'\b(?:\d[ -]*?){13,16}\b')
ip_pat = re.compile(r'\b(?:[0-9]{1,3}\.){3}[0-9]{1,3}\b')
dob_pat = re.compile(r'\b(?:\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}|\d{4}[-/.]\d{1,2}[-/.]\d{1,2})\b')
address_start_pat = re.compile(r'(?i)\b(?:registered office|corporate office|facility at|premises at|leased premises at|manufacturing facility at|plot no\.|village birdewadi|village khalumbre)\b')

all_detections = []

def detect_pii_in_text(text):
    candidates = []
    
    # 1. Email Addresses
    for m in email_pat.finditer(text):
        val = m.group()
        candidates.append((m.start(), m.end(), "email", val, get_fake_email(val)))
        
    # 2. Phone Numbers
    for m in phone_pat.finditer(text):
        val = m.group()
        digits = re.sub(r'\D', '', val)
        if len(digits) >= 10 and len(digits) <= 15:
            candidates.append((m.start(), m.end(), "phone", val, get_fake_phone(val)))
            
    # 3. SSNs
    for m in ssn_pat.finditer(text):
        val = m.group()
        candidates.append((m.start(), m.end(), "ssn", val, "000-00-0000"))
        
    # 4. Credit Cards
    for m in cc_pat.finditer(text):
        val = m.group()
        if luhn_checksum(val):
            candidates.append((m.start(), m.end(), "credit_card", val, "4111-1111-1111-1111"))
            
    # 5. IP Addresses
    for m in ip_pat.finditer(text):
        val = m.group()
        parts = val.split('.')
        if all(0 <= int(x) <= 255 for x in parts):
            candidates.append((m.start(), m.end(), "ip", val, "192.168.1.1"))
            
    # 6. Dates of Birth
    if re.search(r'(?i)\b(?:dob|date of birth|born|d\.o\.b)\b', text):
        for m in dob_pat.finditer(text):
            val = m.group()
            candidates.append((m.start(), m.end(), "dob", val, "01-01-1980"))
            
    # 7. Addresses (Gazetteer & Rules)
    for ad in known_addresses:
        for m in re.finditer(re.escape(ad), text, re.I):
            candidates.append((m.start(), m.end(), "address", m.group(), get_fake_address(ad)))
            
    for m in address_start_pat.finditer(text):
        start = m.start()
        scan_text = text[start:start+180]
        end_match = re.search(r'[;.\n]|\bcs\.connect\b|\bcs\b|\bcompliance\b|\btele\b|\bwebsite\b', scan_text)
        if end_match:
            end = start + end_match.start()
        else:
            end = start + len(scan_text)
            
        addr_val = text[start:end].strip()
        if len(addr_val) > 25 and ("," in addr_val or "Pune" in addr_val or "Maharashtra" in addr_val):
            candidates.append((start, end, "address", addr_val, get_fake_address(addr_val)))

    # 8. Company Names (Gazetteer & Rules)
    for cp in known_companies:
        for m in re.finditer(r'\b' + re.escape(cp) + r'\b', text, re.I):
            val = m.group()
            candidates.append((m.start(), m.end(), "company", val, get_fake_company(val)))
            
    generic_company_pat = re.compile(r'\b[A-Z][a-zA-Z0-9&\s\-\.]+(?:Limited|Private Limited|Pvt\.?\s*Ltd\.?|Ltd\.?|LLP|Corporation|Inc\.)\b')
    for m in generic_company_pat.finditer(text):
        val = m.group()
        if not any(x in val.lower() for x in ["the companies act", "the regional director", "regional director", "central government"]):
            candidates.append((m.start(), m.end(), "company", val, get_fake_company(val)))

    # 9. Names (Gazetteer, NER, and Titles)
    for name in known_names:
        for m in re.finditer(r'\b' + re.escape(name) + r'\b', text, re.I):
            val = m.group()
            candidates.append((m.start(), m.end(), "name", val, get_fake_name(val)))
            
    title_pat = re.compile(r'\b(?:Mr\.|Ms\.|Mrs\.|Shri|Smt\.|Dr\.)\s+[A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,2}\b')
    for m in title_pat.finditer(text):
        val = m.group()
        candidates.append((m.start(), m.end(), "name", val, get_fake_name(val)))

    spacy_doc = nlp(text)
    for ent in spacy_doc.ents:
        if ent.label_ == "PERSON":
            val = ent.text.strip()
            val_clean = re.sub(r'^[^\w\s]+|[^\w\s]+$', '', val).strip()
            if len(val_clean) > 3 and " " in val_clean:
                words = set(val_clean.lower().split())
                if not words.intersection(blacklist_names):
                    candidates.append((ent.start_char, ent.end_char, "name", val, get_fake_name(val)))

    # --- RESOLVE OVERLAPS ---
    candidates = sorted(candidates, key=lambda x: (x[0], -(x[1] - x[0])))
    resolved = []
    last_end = -1
    for start, end, cat, val, fake in candidates:
        if start >= last_end:
            resolved.append((start, end, cat, val, fake))
            last_end = end
            
    return resolved

# --- RUN-LEVEL REPLACE ALGORITHM ---

def replace_in_paragraph(p, matches, location_ref):
    matches = sorted(matches, key=lambda x: x[0], reverse=True)
    
    for start_idx, end_idx, cat, val, fake in matches:
        all_detections.append({
            "location": location_ref,
            "category": cat,
            "original": val,
            "replacement": fake
        })
        
        run_spans = []
        curr_pos = 0
        for r in p.runs:
            run_len = len(r.text)
            run_spans.append((curr_pos, curr_pos + run_len, r))
            curr_pos += run_len
            
        overlapping_runs = []
        for run_start, run_end, r in run_spans:
            if max(run_start, start_idx) < min(run_end, end_idx):
                overlapping_runs.append((run_start, run_end, r))
                
        if not overlapping_runs:
            p.text = p.text[:start_idx] + fake + p.text[end_idx:]
            continue
            
        if len(overlapping_runs) == 1:
            run_start, run_end, r = overlapping_runs[0]
            local_start = start_idx - run_start
            local_end = end_idx - run_start
            r.text = r.text[:local_start] + fake + r.text[local_end:]
        else:
            first_start, first_end, r_first = overlapping_runs[0]
            local_start = start_idx - first_start
            r_first.text = r_first.text[:local_start] + fake
            
            for _, _, r_mid in overlapping_runs[1:-1]:
                r_mid.text = ""
                
            last_start, last_end, r_last = overlapping_runs[-1]
            local_end = end_idx - last_start
            r_last.text = r_last.text[local_end:]

# --- MAIN EXECUTION LOOP ---

def main():
    print("Loading Red Herring Prospectus.docx...")
    doc = docx.Document(doc_path)
    
    # Tracking processed XML element IDs to avoid duplicates from merged cells
    processed_paragraphs = set()
    
    print("Redacting paragraphs...")
    for idx, p in enumerate(doc.paragraphs):
        p_id = id(p._element)
        if p_id in processed_paragraphs:
            continue
        processed_paragraphs.add(p_id)
        
        if not p.text.strip():
            continue
        matches = detect_pii_in_text(p.text)
        if matches:
            replace_in_paragraph(p, matches, f"p[{idx}]")
            
    print("Redacting tables...")
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    p_id = id(p._element)
                    if p_id in processed_paragraphs:
                        continue
                    processed_paragraphs.add(p_id)
                    
                    if not p.text.strip():
                        continue
                    matches = detect_pii_in_text(p.text)
                    if matches:
                        replace_in_paragraph(p, matches, f"t[{t_idx}]r[{r_idx}]c[{c_idx}]p[{p_idx}]")

    print(f"Saving redacted document to {output_path}...")
    doc.save(output_path)
    
    print(f"Writing detections log to {detections_path}...")
    with open(detections_path, "w", encoding="utf-8") as f:
        json.dump(all_detections, f, indent=4)
        
    print(f"Redaction process completed! Mapped {len(all_detections)} PII entities.")

if __name__ == "__main__":
    main()
